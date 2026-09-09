"""Investigation pack Word working copy (INV-PACK-R3).

python-docx from the same stored pack payload the PDF renders. This file is how
a quality lead amends wording *before* issue. INV-C17 still retains PDF bytes
only — Word is not a second disclosure log, and an issued pack must not live
re-render a working copy.
"""

from __future__ import annotations

import io
import re
from typing import Any, Optional

from src.domain.services import investigation_pack_brand as pack_brand
from src.domain.services.investigation_pack_draw import (
    format_stamp,
    humanise_key,
    icam_summary_line,
    normalise_chronology_events,
    normalise_icam_factors,
    plural,
)
from src.domain.services.investigation_pack_pdf import (
    _AUDIENCE_LABELS,
    _CHRONOLOGY_AUDIENCES,
    _CHRONOLOGY_ENTRY_ROWS,
    _CHRONOLOGY_PROVENANCE,
    _CHRONOLOGY_WITHHELD,
    _EMPTY_CAPA,
    _EMPTY_CONTRIBUTING,
    _EMPTY_FINDINGS,
    _EMPTY_ROOT_CAUSE,
    _EMPTY_WHYS,
    _ICAM_HEADING,
    _MAX_ASSET_ROWS,
    _MAX_FIELD_CHARS,
    _PACK_CAPA,
    _PACK_FINDINGS,
    _PACK_ICAM_FACTORS,
    _PACK_ROOT_CAUSE,
    _catalogue_title,
    _human_generated_label,
    _icam_payload,
    _incident_reference,
    chronology_feed,
    confidentiality_notice,
    format_field_value,
    format_pack_field,
    summarise_redactions,
)

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
WORKING_COPY_CLOSED = "PACK_WORKING_COPY_CLOSED"
WORKING_COPY_CLOSED_MESSAGE = "This pack has been issued. The Word working copy is closed; download the retained PDF."
_ICAM_WORD_NOTE = (
    "The ICAM diagram is drawn in the PDF. This Word file lists the same stored factors "
    "so the wording can be edited before issue."
)


def _require_docx() -> Any:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        raise RuntimeError("Word export unavailable: python-docx is not installed in this environment") from exc
    return Document


def _clip(text: str) -> str:
    if len(text) <= _MAX_FIELD_CHARS:
        return text
    return text[: _MAX_FIELD_CHARS - 3] + "..."


_CHAPTER_PREFIX = re.compile(r"^(\d{2})\b")


def _section_bookmark(number: int) -> str:
    return f"pack-sec-{number:02d}"


def _add_bookmark(paragraph: Any, name: str, bookmark_id: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_page_field(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(result)
    run._r.append(end)


def _add_toc_hyperlink(paragraph: Any, text: str, anchor: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "443C38")
    r_pr.append(colour)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    r_pr.append(size)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


class InvestigationPackDocxService:
    """Render a stored investigation customer pack as an editable .docx working copy."""

    @staticmethod
    def docx_filename(investigation_reference: Any, pack_uuid: Any) -> str:
        ref = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in str(investigation_reference or "pack"))
        suffix = str(pack_uuid or "")[:8] or "pack"
        return f"investigation-report-{ref}-{suffix}.docx"

    def build_docx_bytes(  # noqa: C901 - one linear pass matching the PDF
        self,
        pack: dict[str, Any],
        *,
        organisation_name: Optional[str] = None,
        primary_color: Optional[str] = None,
        timeline_events: Any = None,
    ) -> bytes:
        """Render pack bytes. Raises RuntimeError when python-docx is unavailable."""
        _ = organisation_name
        _ = primary_color
        Document = _require_docx()
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor

        raw_content = pack.get("content")
        content: dict[str, Any] = raw_content if isinstance(raw_content, dict) else {}
        audience = str(pack.get("audience") or "")
        audience_label = _AUDIENCE_LABELS.get(audience, humanise_key(audience) or "Customer pack")
        reference = pack.get("investigation_reference") or content.get("investigation_reference") or "Unknown"
        title = pack.get("investigation_title") or content.get("title") or "Investigation report"
        generated_at = pack.get("generated_at") or ""
        confidentiality = confidentiality_notice(audience, pack.get("redaction_log"))
        incident_ref = _incident_reference(pack, content) or str(reference)

        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.different_first_page_header_footer = True

        crimson = RGBColor(*pack_brand.CRIMSON)
        jet = RGBColor(*pack_brand.JET_GREY)

        header = section.header
        header.is_linked_to_previous = False
        header_p = header.paragraphs[0]
        header_run = header_p.add_run(f"INVESTIGATION REPORT  ·  {reference}")
        header_run.font.size = Pt(9)
        header_run.font.color.rgb = jet
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        first_footer = section.first_page_footer
        first_footer.is_linked_to_previous = False
        first_p = first_footer.paragraphs[0]
        first_p.text = pack_brand.legal_footer_line()
        first_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if first_p.runs:
            first_p.runs[0].font.size = Pt(8)
            first_p.runs[0].font.color.rgb = jet

        footer = section.footer
        footer.is_linked_to_previous = False
        footer_p = footer.paragraphs[0]
        left_run = footer_p.add_run(f"{reference}  ·  {pack_brand.CLASSIFICATION}    ")
        left_run.font.size = Pt(8)
        left_run.font.color.rgb = jet
        right_run = footer_p.add_run(f"{audience_label.upper()}  ·  PAGE ")
        right_run.font.size = Pt(8)
        right_run.font.color.rgb = jet
        _add_page_field(footer_p)

        try:
            doc.add_picture(str(pack_brand.lockup_path()), width=Inches(2.6))
        except Exception as exc:  # noqa: BLE001 - missing lockup must fail closed
            raise RuntimeError(f"Investigation pack Word lockup could not be embedded: {exc}") from exc

        eyebrow = doc.add_paragraph()
        run = eyebrow.add_run(f"QUALITY GOVERNANCE PORTAL · {audience_label.upper()}")
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = crimson

        title_p = doc.add_paragraph()
        title_run = title_p.add_run("Investigation Report")
        title_run.bold = True
        title_run.font.size = Pt(28)
        title_run.font.color.rgb = jet

        display = doc.add_paragraph()
        display_run = display.add_run(incident_ref)
        display_run.bold = True
        display_run.font.size = Pt(16)
        display_run.font.color.rgb = crimson

        meta_rows = (
            ("Investigation reference", str(reference)),
            ("Incident reference", incident_ref),
            ("Status", humanise_key(content.get("status") or "unknown")),
            ("Investigation level", humanise_key(content.get("level") or "unknown")),
            ("Pack title", str(title)),
            ("Pack type", audience_label),
            ("Generated", _human_generated_label(generated_at)),
            ("Issued by", pack_brand.ISSUED_BY),
        )
        self._write_kv_table(doc, meta_rows, crimson, jet)

        if confidentiality:
            self._heading(doc, "Confidentiality", crimson, size=12)
            self._body(doc, confidentiality, jet)

        self._body(doc, pack_brand.CLASSIFICATION, jet)

        raw_sections = content.get("sections")
        section_map: dict[str, Any] = raw_sections if isinstance(raw_sections, dict) else {}
        chronology_feed_data, chronology_provenance = chronology_feed(pack, content, timeline_events)
        contents: list[str] = []
        for key in section_map:
            contents.append(_catalogue_title(key))
            fields = section_map[key]
            if str(key) == _PACK_ROOT_CAUSE and isinstance(fields, dict) and _icam_payload(fields) is not None:
                contents.append("Contributing factors")
                contents.append(_ICAM_HEADING)
        if chronology_feed_data is not None:
            contents.append("Chronology")
        contents.extend(("Evidence schedule", "Redaction summary", "Pack integrity"))

        self._heading(doc, "\u2014 CONTENTS", crimson, size=16)
        for index, heading in enumerate(contents, start=1):
            para = doc.add_paragraph()
            _add_toc_hyperlink(para, f"{index} {heading}", _section_bookmark(index))

        chapter = 1
        if not section_map:
            self._body(doc, "No report sections were recorded on this investigation.", jet)
        else:
            for section_key, fields in section_map.items():
                self._heading(doc, f"{chapter:02d} {_catalogue_title(section_key)}", jet, size=14)
                chapter += 1
                if not isinstance(fields, dict) or not fields:
                    self._body(doc, "No content recorded for this section.", jet)
                    continue
                key = str(section_key)
                if key == _PACK_FINDINGS:
                    self._render_findings(doc, fields, jet)
                elif key == _PACK_ROOT_CAUSE:
                    self._render_rca_core(doc, fields, jet)
                    if _icam_payload(fields) is not None:
                        self._heading(doc, f"{chapter:02d} Contributing factors", jet, size=14)
                        chapter += 1
                        self._render_contributing(doc, fields, jet, crimson)
                        self._heading(doc, f"{chapter:02d} {_ICAM_HEADING}", jet, size=14)
                        chapter += 1
                        self._render_icam_list(doc, fields, jet)
                elif key == _PACK_CAPA:
                    self._render_capa(doc, fields, jet, crimson)
                else:
                    rows = tuple((humanise_key(k), format_pack_field(v)) for k, v in fields.items())
                    self._write_kv_table(doc, rows, crimson, jet)

        omitted = content.get("omitted_sections")
        if isinstance(omitted, list) and omitted:
            self._heading(doc, "Sections withheld from this pack", jet, size=14)
            self._body(doc, "The following sections were approved for omission and are not reproduced above:", jet)
            for section_key in omitted:
                self._body(doc, f"- {_catalogue_title(section_key)}", jet)

        if chronology_feed_data is not None:
            self._heading(doc, f"{chapter:02d} Chronology", jet, size=14)
            chapter += 1
            chronology = normalise_chronology_events(chronology_feed_data)
            if not chronology.events:
                self._body(doc, "No chronology events were recorded.", jet)
            elif audience not in _CHRONOLOGY_AUDIENCES:
                self._body(doc, _CHRONOLOGY_WITHHELD, jet)
            else:
                events = list(reversed(chronology.events))[:_CHRONOLOGY_ENTRY_ROWS]
                self._body(doc, f"Most recent entries ({len(events)} of {len(chronology.events)})", jet)
                for event in events:
                    line = f"- {format_stamp(event.at)} - {event.origin_label} - {event.label}"
                    if event.detail:
                        line = f"{line}: {event.detail}"
                    self._body(doc, line, jet)
                if chronology_provenance in _CHRONOLOGY_PROVENANCE:
                    self._body(doc, _CHRONOLOGY_PROVENANCE[chronology_provenance], jet)

        self._heading(doc, f"{chapter:02d} Evidence schedule", jet, size=14)
        chapter += 1
        raw_assets = pack.get("included_assets")
        assets: list[Any] = raw_assets if isinstance(raw_assets, list) else []
        if not assets:
            self._body(doc, "No evidence assets are linked to this investigation.", jet)
        else:
            included = [a for a in assets[:_MAX_ASSET_ROWS] if isinstance(a, dict) and a.get("included")]
            excluded = [a for a in assets[:_MAX_ASSET_ROWS] if isinstance(a, dict) and not a.get("included")]
            self._body(doc, f"Released with this pack: {len(included)}", jet)
            for asset in included:
                self._body(
                    doc,
                    f"- {asset.get('title') or 'Untitled evidence'} ({humanise_key(asset.get('asset_type'))})",
                    jet,
                )
            self._body(doc, f"Withheld: {len(excluded)}", jet)
            for asset in excluded:
                reason = humanise_key(asset.get("exclusion_reason")) if asset.get("exclusion_reason") else "Withheld"
                self._body(doc, f"- {asset.get('title') or 'Untitled evidence'} - {reason}", jet)

        redactions = summarise_redactions(pack.get("redaction_log"))
        self._heading(doc, f"{chapter:02d} Redaction summary", jet, size=14)
        chapter += 1
        if not redactions:
            self._body(doc, "No redactions were applied to this pack.", jet)
        else:
            for kind, count in redactions:
                self._body(doc, f"{humanise_key(kind)}: {count}", jet)

        self._heading(doc, f"{chapter:02d} Pack integrity", jet, size=14)
        self._body(doc, f"Pack UUID: {pack.get('pack_uuid') or 'unknown'}", jet)
        self._body(doc, f"Content SHA-256: {pack.get('checksum_sha256') or 'not recorded'}", jet)
        self._body(
            doc,
            "This Word file is a working copy of the stored pack payload. The issued artefact is the PDF "
            "retained at issue (C17). The SHA-256 above is the checksum of that payload.",
            jet,
        )

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _heading(doc: Any, text: str, colour: Any, *, size: int = 14) -> None:
        from docx.shared import Pt

        para = doc.add_paragraph()
        label = _clip(text)
        if _CHAPTER_PREFIX.match(label.strip()):
            label = label.upper()
        run = para.add_run(label)
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = colour
        match = _CHAPTER_PREFIX.match(label.strip())
        if match:
            number = int(match.group(1))
            _add_bookmark(para, _section_bookmark(number), number)

    @staticmethod
    def _body(doc: Any, text: str, colour: Any) -> None:
        from docx.shared import Pt

        para = doc.add_paragraph()
        run = para.add_run(_clip(text))
        run.font.size = Pt(11)
        run.font.color.rgb = colour

    @staticmethod
    def _write_grid(doc: Any, rows: list[tuple[str, ...]], crimson: Any, jet: Any) -> None:
        from docx.shared import Pt

        if not rows:
            return
        cols = max(len(row) for row in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        for r_index, row in enumerate(rows):
            for c_index in range(cols):
                value = row[c_index] if c_index < len(row) else ""
                cell = table.rows[r_index].cells[c_index]
                cell.text = str(value)
                if not cell.paragraphs[0].runs:
                    continue
                run = cell.paragraphs[0].runs[0]
                run.font.size = Pt(8 if r_index == 0 else 10)
                run.font.color.rgb = crimson if r_index == 0 else jet
                run.bold = r_index == 0

    def _write_kv_table(self, doc: Any, rows: tuple[tuple[str, str], ...], crimson: Any, jet: Any) -> None:
        self._write_grid(doc, [(label, value) for label, value in rows], crimson, jet)

    def _render_findings(self, doc: Any, fields: dict[str, Any], jet: Any) -> None:
        items = fields.get("items")
        if not isinstance(items, list) or not items:
            self._body(doc, _EMPTY_FINDINGS, jet)
            return
        for index, item in enumerate(items, start=1):
            body = item.get("body") if isinstance(item, dict) else item
            self._body(doc, f"{index:02d}  {format_field_value(body)}", jet)

    def _render_why(self, doc: Any, whys: Any, jet: Any) -> None:
        self._heading(doc, "5 Whys", jet, size=12)
        if not isinstance(whys, list) or not whys:
            self._body(doc, _EMPTY_WHYS, jet)
            return
        for raw in whys:
            if not isinstance(raw, dict):
                continue
            level_raw = raw.get("level")
            if level_raw is None:
                continue
            try:
                level = int(level_raw)
            except (TypeError, ValueError):
                continue
            self._heading(doc, f"Why {level}", jet, size=11)
            question = raw.get("why")
            if isinstance(question, str) and question.strip():
                self._body(doc, f"Question: {question.strip()}", jet)
            self._body(doc, f"Answer: {format_pack_field(raw.get('answer'))}", jet)
            evidence = raw.get("evidence")
            if isinstance(evidence, str) and evidence.strip():
                self._body(doc, f"Evidence: {evidence.strip()}", jet)

    def _render_rca_core(self, doc: Any, fields: dict[str, Any], jet: Any) -> None:
        self._stated(
            doc, "Problem statement", fields.get("problem_statement"), "No problem statement was recorded.", jet
        )
        self._render_why(doc, fields.get("whys"), jet)
        self._stated(doc, "Root cause", fields.get("root_cause"), _EMPTY_ROOT_CAUSE, jet)
        if _icam_payload(fields) is None:
            self._stated(doc, "Contributing factors", fields.get("contributing_factors"), _EMPTY_CONTRIBUTING, jet)

    def _render_contributing(self, doc: Any, fields: dict[str, Any], jet: Any, crimson: Any) -> None:
        raw = _icam_payload(fields)
        factors = normalise_icam_factors(raw)
        if factors.factors:
            rows: list[tuple[str, ...]] = [("ICAM category", "Contributing factor", "HSG245 causal depth")]
            for factor in factors.factors:
                if factor.sub_causes:
                    bullets = "; ".join(factor.sub_causes)
                    body = f"{factor.cause} ({bullets})"
                else:
                    body = factor.cause
                rows.append((factor.category_label, body, factor.depth_label or "—"))
            self._write_grid(doc, rows, crimson, jet)
            return
        self._stated(doc, "Contributing factors", fields.get("contributing_factors"), _EMPTY_CONTRIBUTING, jet)

    def _render_icam_list(self, doc: Any, fields: dict[str, Any], jet: Any) -> None:
        raw = fields.get(_PACK_ICAM_FACTORS)
        if raw is None:
            return
        factors = normalise_icam_factors(raw)
        self._body(doc, icam_summary_line(factors), jet)
        self._body(doc, _ICAM_WORD_NOTE, jet)
        for factor in factors.factors:
            line = f"- {factor.category_label}: {factor.label}"
            if factor.depth_label:
                line = f"{line} [{factor.depth_label.lower()}]"
            self._body(doc, line, jet)
        if factors.omitted:
            self._body(
                doc,
                f"{plural(factors.omitted, 'further factor', 'further factors')} "
                f"{'is' if factors.omitted == 1 else 'are'} not shown in the diagram, which is capped at the "
                f"first {len(factors.factors)} in ICAM category order.",
                jet,
            )

    def _stated(self, doc: Any, heading: str, value: Any, empty_message: str, jet: Any) -> None:
        self._heading(doc, heading, jet, size=12)
        if isinstance(value, str) and not value.strip():
            self._body(doc, empty_message, jet)
            return
        if isinstance(value, list) and not value:
            self._body(doc, empty_message, jet)
            return
        self._body(doc, format_field_value(value), jet)

    def _render_capa(self, doc: Any, fields: dict[str, Any], jet: Any, crimson: Any) -> None:
        items = fields.get("items")
        if not isinstance(items, list) or not items:
            self._body(doc, _EMPTY_CAPA, jet)
            return
        rows: list[tuple[str, str]] = []
        for item in items:
            if not isinstance(item, dict):
                rows.append(("", format_pack_field(item)))
                continue
            title = str(item.get("title") or "").strip() or "—"
            reference = str(item.get("reference") or "").strip() or "—"
            why_level = item.get("why_level")
            action = title
            if why_level is not None:
                try:
                    action = f"{action} (Why {int(why_level)})"
                except (TypeError, ValueError):
                    pass
            rows.append((reference, action))
        self._write_grid(doc, [("Reference", "Action"), *rows], crimson, jet)

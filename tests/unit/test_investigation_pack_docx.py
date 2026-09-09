"""INV-PACK-R3: Word working copy from the same stored pack payload."""

from __future__ import annotations

import builtins
import io

import pytest

from src.domain.services.investigation_pack_docx import InvestigationPackDocxService
from src.domain.services.investigation_pack_pdf import InvestigationPackPdfService


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [para.text for para in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _pack(**overrides) -> dict:
    pack = {
        "pack_uuid": "6f1c2d3e-0000-4000-8000-abcdefabcdef",
        "audience": "external_customer",
        "investigation_reference": "INV-2026-0007",
        "investigation_title": "Collision on the A1",
        "generated_at": "2026-07-20T10:00:00+00:00",
        "checksum_sha256": "a" * 64,
        "content": {
            "investigation_reference": "INV-2026-0007",
            "title": "Collision on the A1",
            "status": "completed",
            "level": "high",
            "sections": {
                "section_1_details": {
                    "incident_date": "2026-05-17",
                    "site_conditions": "Wet road surface, poor visibility",
                },
                "findings": {"items": [{"body": "Guard was missing from the mill"}]},
                "root-cause": {
                    "whys": [{"level": 1, "why": "", "answer": "The guard had been removed"}],
                    "root_cause": "No permit",
                    "contributing_factors": "",
                },
                "capa": {
                    "items": [
                        {"title": "Replace the guard", "reference": "CAPA-2026-0042", "why_level": 1},
                    ]
                },
            },
            "omitted_sections": ["section_5_internal_commentary"],
        },
        "included_assets": [
            {"asset_id": 1, "title": "Dashcam still", "asset_type": "photo", "included": True},
        ],
    }
    pack.update(overrides)
    return pack


def test_renders_an_ooxml_document() -> None:
    out = InvestigationPackDocxService().build_docx_bytes(_pack())

    assert out.startswith(b"PK")
    assert InvestigationPackDocxService.docx_filename("INV-2026-0007", "6f1c2d3e-0000") == (
        "investigation-report-INV-2026-0007-6f1c2d3e.docx"
    )


def test_contents_and_body_match_the_pdf_catalogue() -> None:
    text = _docx_text(InvestigationPackDocxService().build_docx_bytes(_pack()))

    assert "01 Incident details" in text
    assert "02 Findings" in text
    assert "17 May 2026" in text
    assert "01  Guard was missing from the mill" in text
    assert "Why 1" in text
    assert "The guard had been removed" in text
    assert "Why: Not recorded" not in text
    assert "CAPA-2026-0042" in text
    assert "Replace the guard (Why 1)" in text
    assert "Plantexpand Ltd" in text
    assert "UNCONTROLLED WHEN PRINTED" in text
    assert "This Word file is a working copy" in text


def _timeline_events() -> list[dict]:
    return [
        {
            "id": -111,
            "event_type": "SOURCE_AUDIT",
            "new_value": "Incident raised",
            "actor_name": "Dana Reporter",
            "event_metadata": {"origin": "source", "source_label": "Incident · create"},
            "created_at": "2026-05-01T08:00:00+00:00",
        }
    ]


def test_external_pack_still_withholds_chronology() -> None:
    text = _docx_text(InvestigationPackDocxService().build_docx_bytes(_pack(), timeline_events=_timeline_events()))

    assert "The chronology is withheld from this pack." in text
    assert "Dana Reporter" not in text


def test_icam_factors_are_listed_not_drawn() -> None:
    pack = _pack()
    pack["content"]["sections"]["root-cause"]["icam_factors"] = {
        "factors": [
            {
                "id": 4,
                "category": "organisational_factors",
                "cause": "No refresher training schedule",
                "sub_causes": ["Budget withdrawn"],
                "depth": "underlying",
            }
        ],
        "unmapped_categories": [],
        "unpresentable": 0,
    }
    text = _docx_text(InvestigationPackDocxService().build_docx_bytes(pack))

    assert "ICAM contributing factors" in text
    assert "No refresher training schedule" in text
    assert "The ICAM diagram is drawn in the PDF" in text
    assert "No refresher training schedule..." not in text


def test_fails_closed_when_the_lockup_cannot_be_embedded(monkeypatch: pytest.MonkeyPatch) -> None:
    from pathlib import Path

    from src.domain.services import investigation_pack_brand as pack_brand

    monkeypatch.setattr(pack_brand, "lockup_path", lambda: Path("/tmp/missing-lockup.png"))

    with pytest.raises(RuntimeError, match="lockup could not be embedded"):
        InvestigationPackDocxService().build_docx_bytes(_pack())


def test_tenant_branding_is_ignored() -> None:
    text = _docx_text(
        InvestigationPackDocxService().build_docx_bytes(
            _pack(), organisation_name="Default Organisation", primary_color="#3B82F6"
        )
    )

    assert "Default Organisation" not in text
    assert "Plantexpand Ltd" in text


def test_fails_closed_when_python_docx_is_missing(monkeypatch: pytest.MonkeyPatch) -> None:
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "docx" or name.startswith("docx."):
            raise ModuleNotFoundError("No module named 'docx'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)

    with pytest.raises(RuntimeError, match="python-docx is not installed"):
        InvestigationPackDocxService().build_docx_bytes(_pack())


def test_pdf_and_word_filenames_share_the_same_stem() -> None:
    pdf = InvestigationPackPdfService.pdf_filename("INV-2026-0007", "6f1c2d3e-0000")
    docx = InvestigationPackDocxService.docx_filename("INV-2026-0007", "6f1c2d3e-0000")

    assert pdf.replace(".pdf", "") == docx.replace(".docx", "")


def test_long_capa_titles_are_not_clipped_in_word() -> None:
    pack = _pack()
    pack["content"]["sections"]["capa"] = {
        "items": [
            {
                "title": "Containment: restrict S Leggitt from chainsaw work on Forestry England sites",
                "reference": "CAPA-2026-0010",
            }
        ]
    }
    text = _docx_text(InvestigationPackDocxService().build_docx_bytes(pack))

    assert "Forestry England sites" in text
    assert "CAPA-2026-0010" in text


def test_contents_rows_hyperlink_to_section_bookmarks() -> None:
    from docx import Document

    document = Document(io.BytesIO(InvestigationPackDocxService().build_docx_bytes(_pack())))
    body = document.element.body.xml
    assert 'w:anchor="pack-sec-01"' in body
    assert 'w:name="pack-sec-01"' in body
    assert 'w:name="pack-sec-02"' in body
